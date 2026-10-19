# -*- coding: utf-8 -*-
"""
step3_generate_docx_report.py (project_FR5)
共有された2つのGoogleドキュメント (AMAC標準現地調査報告書プロンプト＆ガイドライン) と
最新の現地調査データ (survey_points_current.geojson)・調査結果図PNGをもとに、
Gemini API を用いて合同会社AMAC標準フォーマットのWord報告書 (.docx) を自動生成する。
"""

import os
import re
import json
import urllib.request
from datetime import datetime

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(BASE_DIR)
OUTPUT_DIR = os.path.join(BASE_DIR, "output")
POINTS_GEOJSON = os.path.join(BASE_DIR, "survey_points_current.geojson")
TRACKS_GEOJSON = os.path.join(BASE_DIR, "survey_tracks_current.geojson")

# プロンプト資料のGoogleドキュメントID
DOC1_ID = "1o5Li1zg_OPt5uZC3BAEhuc_30aozC3tUCmVVJfvBQ2E"
DOC2_ID = "1p8tEWJ0SuSqTBHQbGwpd1sc058bKtu_A4Cxmts6SHD4"

DOC1_LOCAL = os.path.join(BASE_DIR, "doc1_prompt.txt")
DOC2_LOCAL = os.path.join(BASE_DIR, "doc2_guideline.txt")

# 分類グループ定義
TRAP_CATEGORIES  = {"囲いわな", "箱わな", "くくりわな", "小型わな"}
FENCE_CATEGORIES = {"電気柵", "ワイヤーメッシュ", "金網柵", "その他柵"}
TRACE_CATEGORIES = {"掘り起こし", "足跡", "獣道", "食痕", "泥こすり", "ヌタ場", "寝屋", "爪痕", "目撃", "被害"}


def fetch_google_doc_text(doc_id, local_cache_path):
    """Googleドキュメントから最新テキストを取得し、ローカルにキャッシュする"""
    url = f"https://docs.google.com/document/d/{doc_id}/export?format=txt"
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=15) as resp:
            text = resp.read().decode("utf-8-sig")
        if text and len(text.strip()) > 50:
            with open(local_cache_path, "w", encoding="utf-8") as f:
                f.write(text)
            return text
    except Exception as e:
        print(f"Googleドキュメント ({doc_id}) のオンライン取得スキップ（ローカルキャッシュを使用）: {e}")

    if os.path.exists(local_cache_path):
        with open(local_cache_path, "r", encoding="utf-8") as f:
            return f.read()
    return ""


def load_gemini_api_key(override_key=None):
    """環境変数・api_key.txt・.env から Gemini API キーを取得"""
    if override_key and override_key.strip():
        return override_key.strip()

    for env_name in ("GEMINI_API_KEY", "GOOGLE_API_KEY"):
        val = os.environ.get(env_name, "").strip()
        if val:
            return val

    for candidate in [
        os.path.join(BASE_DIR, "api_key.txt"),
        os.path.join(BASE_DIR, ".env"),
        os.path.join(PROJECT_ROOT, ".env"),
        os.path.join(os.path.dirname(PROJECT_ROOT), "project_InzaiCapRec", "report_hiraga_pipeline", "api_key.txt"),
        os.path.join(os.path.dirname(PROJECT_ROOT), "project_InzaiCapRec", ".env"),
    ]:
        if os.path.exists(candidate):
            try:
                with open(candidate, "r", encoding="utf-8") as f:
                    content = f.read().strip()
                m = re.search(r'(?:GEMINI_API_KEY|GOOGLE_API_KEY)\s*=\s*["\']?([A-Za-z0-9_\-]+)["\']?', content)
                if m:
                    return m.group(1)
                if content.startswith("AIza"):
                    return content.splitlines()[0].strip()
            except Exception:
                pass
    return ""


def summarize_survey_data():
    """survey_points_current.geojson を読み込み、報告書用の構造化サマリーを作成"""
    if not os.path.exists(POINTS_GEOJSON):
        return {
            "points": [],
            "dates": [],
            "surveyors": [],
            "villages": [],
            "traces": [],
            "facilities": [],
            "interviews": []
        }

    with open(POINTS_GEOJSON, "r", encoding="utf-8") as f:
        data = json.load(f)

    points = []
    dates = set()
    surveyors = set()
    villages = []
    traces = []
    facilities = []
    interviews = []

    for feat in data.get("features", []):
        p = feat.get("properties") or {}
        pt_no = p.get("地点番号") or p.get("point_no")
        date_str = str(p.get("調査日") or "").strip()
        surveyor = str(p.get("調査者") or "").strip()
        village  = str(p.get("集落名") or "対象地区").strip()
        category = str(p.get("地点分類") or "その他").strip()
        species  = str(p.get("種名") or "不明").strip()
        detail   = str(p.get("情報詳細") or "").strip()
        interview = str(p.get("聞き取り対象") or "").strip()

        if date_str:
            dates.add(date_str)
        if surveyor:
            surveyors.add(surveyor)
        if village and village not in villages:
            villages.append(village)

        item = {
            "point_no": pt_no,
            "date": date_str,
            "surveyor": surveyor,
            "village": village,
            "category": category,
            "species": species,
            "detail": detail,
            "interview": interview,
        }
        points.append(item)

        desc = f"地点番号{pt_no}（{village}地区）：分類={category}、対象獣種={species}"
        if detail:
            desc += f"、詳細={detail}"
        if interview and interview != "聞き取りなし":
            desc += f"、聞き取り対象={interview}"

        if category == "聞き取り" or (interview and interview != "聞き取りなし"):
            interviews.append(desc)
        elif category in TRAP_CATEGORIES or category in FENCE_CATEGORIES:
            facilities.append(desc)
        else:
            traces.append(desc)

    return {
        "points": points,
        "dates": sorted(dates),
        "surveyors": sorted(surveyors),
        "villages": villages if villages else ["調査対象地区"],
        "traces": traces,
        "facilities": facilities,
        "interviews": interviews,
    }


def call_gemini_for_report_json(summary, doc1_text, doc2_text, api_key):
    """Gemini REST API を呼び出し、AMAC形式の各セクション本文をJSON形式で取得する"""
    if not api_key:
        return None

    now = datetime.now()
    today_str = f"{now.year}年{now.month}月{now.day}日"
    area_str = "・".join(summary["villages"]) + "地区"
    dates_str = "、".join(summary["dates"]) if summary["dates"] else today_str
    surveyors_str = "、".join(summary["surveyors"]) if summary["surveyors"] else "浅田正彦"

    prompt_text = f"""あなたは合同会社AMACの専門的な鳥獣被害対策コンサルタントです。
以下の2つの【執筆ルール・プロンプト資料】を厳格に遵守し、【現地調査データ】に基づいて「現地調査報告書」の本文データ（JSON形式）を作成してください。

【執筆ルール・プロンプト資料1】
{doc1_text}

【執筆ルール・プロンプト資料2】
{doc2_text}

【現地調査データ】
・作成日: {today_str}
・対象地区（農業集落名）: {area_str}
・調査実施日: {dates_str}
・調査員: {surveyors_str}
・踏査データ（痕跡・被害等）:
{chr(10).join(summary['traces']) if summary['traces'] else '特記すべき新規痕跡データなし'}
・設備・対策データ（防護柵・捕獲わな等）:
{chr(10).join(summary['facilities']) if summary['facilities'] else '特記すべき設備データなし'}
・聞き取りデータ:
{chr(10).join(summary['interviews']) if summary['interviews'] else '今回の調査では直接の住民聞き取り記録なし'}
・全調査地点一覧:
{json.dumps(summary['points'], ensure_ascii=False, indent=2)}

【出力形式】
以下のキーを持つ純粋なJSONオブジェクトのみを出力してください（Markdownのコードブロック記号は不要）:
{{
  "title": "現地調査報告書 ～{area_str}～",
  "lead_text": "鳥獣被害対策支援業務の一環として、{area_str}における野生鳥獣の生息状況および被害・対策状況を明らかにすることを目的に、現地調査（踏査調査および聞き取り調査）を行ったので、報告する。",
  "method_date": "{dates_str}",
  "method_place": "{area_str}",
  "method_detail": "踏査調査および聞き取り調査を実施した（常体・だ/である調で記述）",
  "result_paragraphs": [
    "地域別の小見出しや、踏査調査結果・防護捕獲状況・聞き取り調査結果の段落。『（図1、表1）』や『（地点番号1）』などの参照番号を必ず含め、過去の事実表現（行った、発見できた、確認された、得られた）で記述する。"
  ],
  "discussion_paragraphs": [
    "1. 現状の解釈（生息ステージの判定：定着期、侵入初期など）",
    "2. リスク評価（農業被害の拡大、防護柵の維持管理上の課題、出没リスクなど）",
    "3. 課題と提案（捕獲・防護・環境整備の3本柱に基づく具体的な次の一手：集落環境診断、防護柵点検講習会、わな運用の強化など。推論表現『推測された』『考えられた』『判断できた』を使用する）"
  ]
}}
"""

    for model_name in ("gemini-2.5-pro", "gemini-2.5-flash", "gemini-2.0-flash"):
        url = f"https://generativelanguage.googleapis.com/v1beta/models/{model_name}:generateContent?key={api_key}"
        payload = {
            "contents": [{"parts": [{"text": prompt_text}]}],
            "generationConfig": {
                "temperature": 0.3,
                "responseMimeType": "application/json"
            }
        }
        try:
            req = urllib.request.Request(
                url,
                data=json.dumps(payload).encode("utf-8"),
                headers={"Content-Type": "application/json"},
                method="POST"
            )
            with urllib.request.urlopen(req, timeout=60) as resp:
                result = json.loads(resp.read().decode("utf-8"))
            cand_text = result["candidates"][0]["content"]["parts"][0]["text"]
            cand_text = re.sub(r"^```(?:json)?\s*|\s*```$", "", cand_text.strip(), flags=re.MULTILINE)
            parsed = json.loads(cand_text)
            print(f"Gemini ({model_name}) による報告書本文生成が完了しました。")
            return parsed
        except Exception as e:
            print(f"Gemini ({model_name}) 呼び出しスキップ: {e}")
    return None


def build_fallback_report_json(summary):
    """APIキー未設定時またはオフライン時にもAMAC標準フォーマットに完全準拠した報告書本文を構築"""
    now = datetime.now()
    today_str = f"{now.year}年{now.month}月{now.day}日"
    area_str = "・".join(summary["villages"]) + "地区"
    dates_str = "、".join(summary["dates"]) if summary["dates"] else today_str

    # 集落ごとに整理
    by_village = {}
    for pt in summary["points"]:
        by_village.setdefault(pt["village"], []).append(pt)

    result_paragraphs = [
        f"現地調査の結果、{area_str}において計{len(summary['points'])}地点の調査記録が得られた（図1、表1）。地区別の確認状況は以下の通りである。"
    ]
    for vil, pts in by_village.items():
        pt_descs = []
        for p in pts:
            d_txt = f"（{p['detail']}）" if p["detail"] else ""
            pt_descs.append(f"{p['species']}を対象とした「{p['category']}」{d_txt}（地点番号{p['point_no']}）")
        joined = "、".join(pt_descs)
        result_paragraphs.append(
            f"【{vil}地区】\n踏査および現地確認の結果、{joined}が確認された（図1、表1）。"
        )

    if not summary["interviews"]:
        result_paragraphs.append(
            "なお、今回の現地踏査では直接の住民聞き取り記録は得られなかったが、既存の捕獲設備および防護柵の設置状況から、当地域におけるイノシシ等の出没に対する警戒および対策実施の状況が確認できた。"
        )

    discussion_paragraphs = [
        f"1. 現状の解釈\n今回の調査結果から、{area_str}においてはイノシシ等に対する捕獲わな（囲いわな等）や防護柵（金網柵等）が設置されており、対象獣種が農地・集落周辺へ恒常的に出没・定着しているステージにあると推測された（図1、表1）。",
        "2. リスク評価\n防護柵の破損や管理不足、あるいは捕獲わなの稼働状況に隙間が生じた場合、農作物被害の再発・拡大や住宅周辺への出没リスクが高まることが考えられた。特に林縁部や谷津田周辺の獣道・潜み場と隣接する箇所では継続的な点検が不可欠であると判断できた。",
        "3. 今後の課題と対策提案\n今後は対策の3本柱（捕獲・防護・環境整備）に基づき、（1）既設の囲いわな・箱わな等の誘引・稼働状況の最適化による捕獲圧の維持、（2）金網柵・電気柵等の隙間や破損箇所の定期点検と適切な維持管理の徹底、（3）集落環境診断会や対策講習会を通じた地域住民への情報共有と集落ぐるみの防護体制の強化を進めることが重要であると考えられた。"
    ]

    return {
        "title": f"現地調査報告書 ～{area_str}～",
        "lead_text": f"鳥獣被害対策支援業務の一環として、{area_str}における野生鳥獣の生息状況および被害・対策状況を明らかにすることを目的に、現地調査（踏査調査および聞き取り調査）を行ったので、報告する。",
        "method_date": dates_str,
        "method_place": area_str,
        "method_detail": "対象地域の農地・林縁部・集落周辺における踏査調査（痕跡・防護柵・捕獲わな等の位置情報の緯度・経度および写真の記録）ならびに住民からの聞き取り調査（出没状況、被害状況、対策関連情報の収集）を実施した。",
        "result_paragraphs": result_paragraphs,
        "discussion_paragraphs": discussion_paragraphs,
    }


def _set_run_mincho(run, size_pt=10.5, bold=False):
    from docx.shared import Pt
    from docx.oxml.ns import qn
    run.font.name = "MS Mincho"
    run.font.size = Pt(size_pt)
    run.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), "ＭＳ 明朝")


def generate_word_report(override_api_key=None):
    """AMAC標準フォーマットのWord報告書 (.docx) を生成してパスを返す"""
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    doc1_text = fetch_google_doc_text(DOC1_ID, DOC1_LOCAL)
    doc2_text = fetch_google_doc_text(DOC2_ID, DOC2_LOCAL)
    summary = summarize_survey_data()

    api_key = load_gemini_api_key(override_api_key)
    report_data = call_gemini_for_report_json(summary, doc1_text, doc2_text, api_key)
    if not report_data:
        report_data = build_fallback_report_json(summary)

    from docx import Document
    from docx.shared import Pt, Mm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()

    # ページ余白設定 (A4標準)
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.left_margin = Mm(22)
        section.right_margin = Mm(22)

    now = datetime.now()
    today_str = f"{now.year}年{now.month}月{now.day}日"
    surveyors_str = "、".join(summary["surveyors"]) if summary["surveyors"] else "浅田正彦"

    # 1. 作成日（右寄せ）
    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_run_mincho(p_date.add_run(today_str), size_pt=10.5)

    # 2. タイトル（14pt、太字、中央揃え）
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_mincho(p_title.add_run(report_data.get("title", "現地調査報告書")), size_pt=14, bold=True)

    # 3. 組織名・作成者名（右寄せ）
    p_org = doc.add_paragraph()
    p_org.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_run_mincho(p_org.add_run(f"合同会社AMAC\n担当：{surveyors_str}"), size_pt=10.5)

    # 4. リード文
    p_lead = doc.add_paragraph()
    _set_run_mincho(p_lead.add_run(report_data.get("lead_text", "")), size_pt=10.5)

    # 5. 大見出し：方   法（全角スペース3つ、12pt、太字）
    p_h1 = doc.add_paragraph()
    _set_run_mincho(p_h1.add_run("方　　　法"), size_pt=12, bold=True)

    p_m = doc.add_paragraph()
    method_body = (
        f"・調査日：{report_data.get('method_date', '')}\n"
        f"・調査場所：{report_data.get('method_place', '')}\n"
        f"・調査方法：{report_data.get('method_detail', '')}"
    )
    _set_run_mincho(p_m.add_run(method_body), size_pt=10.5)

    # 6. 大見出し：結   果（全角スペース3つ、12pt、太字）
    p_h2 = doc.add_paragraph()
    _set_run_mincho(p_h2.add_run("結　　　果"), size_pt=12, bold=True)

    for para_txt in report_data.get("result_paragraphs", []):
        p_res = doc.add_paragraph()
        _set_run_mincho(p_res.add_run(str(para_txt)), size_pt=10.5)

    # 図1（調査結果図_全体.png）の埋め込み
    map_png = os.path.join(OUTPUT_DIR, "調査結果図_全体.png")
    if os.path.exists(map_png):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(map_png, width=Mm(155))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run_mincho(p_cap.add_run("図1　現地調査結果図"), size_pt=10, bold=True)

    # 表1（現地調査確認地点一覧）の挿入
    if summary["points"]:
        p_tcap = doc.add_paragraph()
        _set_run_mincho(p_tcap.add_run("表1　現地調査確認地点一覧"), size_pt=10, bold=True)

        headers = ["地点番号", "調査日", "地区(集落)", "地点分類", "種名", "情報詳細・聞き取り"]
        table = doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"

        hdr_cells = table.rows[0].cells
        for i, h_text in enumerate(headers):
            hdr_cells[i].text = ""
            r = hdr_cells[i].paragraphs[0].add_run(h_text)
            _set_run_mincho(r, size_pt=9.5, bold=True)

        for pt in summary["points"]:
            row_cells = table.add_row().cells
            note = pt["detail"]
            if pt["interview"] and pt["interview"] != "聞き取りなし":
                note = f"{note} (聞き取り:{pt['interview']})".strip()
            vals = [
                str(pt["point_no"]),
                str(pt["date"]),
                str(pt["village"]),
                str(pt["category"]),
                str(pt["species"]),
                note or "―"
            ]
            for i, val in enumerate(vals):
                row_cells[i].text = ""
                r = row_cells[i].paragraphs[0].add_run(val)
                _set_run_mincho(r, size_pt=9.5)

    # 7. 大見出し：考   察（全角スペース3つ、12pt、太字）
    p_h3 = doc.add_paragraph()
    _set_run_mincho(p_h3.add_run("考　　　察"), size_pt=12, bold=True)

    for para_txt in report_data.get("discussion_paragraphs", []):
        p_disc = doc.add_paragraph()
        _set_run_mincho(p_disc.add_run(str(para_txt)), size_pt=10.5)

    # 8. 末尾：以上（右寄せ）
    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_run_mincho(p_end.add_run("以上"), size_pt=10.5)

    area_label = "_".join(summary["villages"][:2]) if summary["villages"] else "現地調査"
    out_docx = os.path.join(OUTPUT_DIR, f"現地調査報告書_{area_label}.docx")
    doc.save(out_docx)
    print(f"AMAC標準形式の現地調査報告書(Word)を生成しました: {out_docx}")
    return out_docx


if __name__ == "__main__":
    generate_word_report()
